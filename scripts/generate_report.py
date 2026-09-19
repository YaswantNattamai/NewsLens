"""
Generate a Word (.docx) project report for NewsLens, suitable for showing to a
project guide / advisor.

Pulls a live worked-example from the running API if it's up (falls back to
recorded values otherwise), so the numbers in the report are real.

Run:  python -m scripts.generate_report
Output: NewsLens_Project_Report.docx in the repo root.
"""
from __future__ import annotations

import json
import urllib.request

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, RGBColor, Inches

# ---------------------------------------------------------------------------
# Palette / helpers
# ---------------------------------------------------------------------------
INK = RGBColor(0x1C, 0x1B, 0x19)
BLUE = RGBColor(0x2B, 0x45, 0x70)
RED = RGBColor(0xA6, 0x3D, 0x40)
SOFT = RGBColor(0x6E, 0x6A, 0x61)

API = "http://localhost:8000"


def _get(path):
    return json.load(urllib.request.urlopen(API + path, timeout=5))


def fetch_example():
    """Real worked example from the running API; recorded fallback if it's down."""
    fallback = {
        "sources": ["fox", "hpo", "nyt"],
        "n_entities": 129,
        "by_type": {"PERSON": 44, "ORG": 54, "GPE": 31},
        "framing": [("fox", "hpo", 40, 17), ("fox", "nyt", 36, 17), ("hpo", "nyt", 39, 20)],
        "omission": {"fox": 56, "hpo": 38, "nyt": 53},
        "pair": {
            "sim": 0.50,
            "a": "The Associated Press reported that McGurk said in a resignation letter to "
                 "Secretary of State Mike Pompeo that ISIS was on the run, but wasn't yet "
                 "defeated.",
            "b": "Additionally, forces were withdrawn despite the concern by many Republicans "
                 "that leaving would strengthen the hand of Russia and Iran, which both "
                 "support Syria's government.",
        },
        "live": False,
    }
    try:
        arts = _get("/events/2/articles")
        ents = _get("/events/2/entities")
        fr = _get("/events/2/framing")
        om = _get("/events/2/omission")
        from collections import Counter
        by_type = dict(Counter(e["entity_type"] for e in ents))
        framing = [(r["source_a"], r["source_b"], round((r["framing_score"] or 0) * 100),
                    len(r["aligned_pairs"] or [])) for r in fr]
        seen = {}
        for r in om:
            seen.setdefault(r["source_a"], round((r["omission_score_a"] or 0) * 100))
            seen.setdefault(r["source_b"], round((r["omission_score_b"] or 0) * 100))
        r0 = fr[0]
        ap = sorted(r0["aligned_pairs"], key=lambda p: p["similarity"])[0]
        return {
            "sources": [a["source"] for a in arts],
            "n_entities": len(ents),
            "by_type": by_type,
            "framing": framing,
            "omission": seen,
            "pair": {"sim": ap["similarity"], "a": ap["sentence_a"], "b": ap["sentence_b"]},
            "live": True,
        }
    except Exception:
        return fallback


# ---------------------------------------------------------------------------
# Document building helpers
# ---------------------------------------------------------------------------
def add_heading(doc, text, level, color=BLUE):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = color
    return h


def body(doc, text, size=11, italic=False, color=INK, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.italic = italic
    run.font.color.rgb = color
    return p


def bullet(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
        r.font.color.rgb = INK
    r2 = p.add_run(text)
    r2.font.color.rgb = INK
    return p


def styled_table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(10)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(10)
    return t


# ---------------------------------------------------------------------------
# Build
# ---------------------------------------------------------------------------
def build():
    ex = fetch_example()
    doc = Document()

    # default font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # ---------- Cover ----------
    doc.add_paragraph().add_run("\n\n")
    eyebrow = doc.add_paragraph()
    eyebrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    er = eyebrow.add_run("NARRATIVE BIAS DESK  ·  PROJECT REPORT")
    er.font.size = Pt(11)
    er.font.color.rgb = SOFT
    er.bold = True

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("NewsLens")
    tr.font.size = Pt(40)
    tr.bold = True
    tr.font.color.rgb = BLUE

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("Detecting how news sources differ in covering the same story —\n"
                     "entities, sentiment, framing, and omission across the political spectrum")
    sr.font.size = Pt(13)
    sr.italic = True
    sr.font.color.rgb = INK

    doc.add_paragraph().add_run("\n")
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run("An NLP system for computational media-bias analysis\n"
                      "Built on the BASIL dataset, extended with live news ingestion and metric validation\n\n"
                      "Prepared: 13 July 2026")
    mr.font.size = Pt(11)
    mr.font.color.rgb = SOFT

    doc.add_page_break()

    # ---------- Abstract ----------
    add_heading(doc, "Abstract", 1)
    body(doc,
         "NewsLens is a natural-language-processing system that reveals how different news "
         "outlets cover the same event differently. Given a set of articles on one story, it "
         "extracts and canonicalizes named entities across sources, estimates each source's "
         "sentiment toward those entities, quantifies how divergently sources frame comparable "
         "sentences, measures what each source omits relative to the others, and groups the "
         "coverage into narrative threads. The core pipeline is built on five NLP models "
         "(four pretrained, one optionally fine-tuned) and is served through a FastAPI backend "
         "and a React frontend.")
    body(doc,
         "This report documents the completed system and two extensions that move it from a "
         "research demo toward a genuinely useful public tool: (1) live ingestion, which lets a "
         "user analyze how outlets are covering any current topic by pulling real-time coverage "
         "from the open GDELT news index; and (2) a trust-and-validation layer, which measures "
         "the framing metric against human bias annotations and reframes the entire interface "
         "around presenting differences rather than verdicts. A key, deliberately honest finding "
         "is that the framing metric correlates only weakly with human bias labels "
         "(ROC AUC ≈ 0.56), which the system now communicates explicitly to users.")

    # ---------- 1. Introduction ----------
    add_heading(doc, "1. Introduction and Motivation", 1)
    body(doc,
         "Readers rarely see how the outlet they trust differs from others covering the same "
         "event. Two newspapers can report identical facts yet leave a reader with opposite "
         "impressions through word choice, which details they emphasize, whom they quote, and "
         "what they leave unsaid. These effects — framing and omission — are subtle, cumulative, "
         "and hard to notice from inside a single source.")
    body(doc,
         "NewsLens makes these differences visible and side-by-side. Rather than labeling any "
         "outlet 'biased', it surfaces measurable signals of divergence and invites the reader to "
         "compare sources directly. The design goal throughout is honesty: every number is "
         "presented as a signal of difference, bounded by its known limitations, never as a "
         "verdict on truth.")
    body(doc, "The project answers three research questions:")
    bullet(doc, "How do sources differ in the entities they name and the sentiment they express toward them?", "RQ1 — Entities & Sentiment: ")
    bullet(doc, "How divergently do sources frame the same underlying details, and does that divergence track human-perceived bias?", "RQ2 — Framing: ")
    bullet(doc, "What does each source omit relative to the others, at both the entity and thematic level?", "RQ3 — Omission: ")

    # ---------- 2. Architecture ----------
    add_heading(doc, "2. System Architecture", 1)
    body(doc,
         "A single request (POST /events for a curated event, or POST /live-events for a live "
         "topic) drives the whole pipeline. Articles are loaded and stored, then a per-event "
         "orchestrator runs each analysis stage and persists the results to PostgreSQL, which the "
         "API exposes to the React frontend.")
    body(doc, "End-to-end flow:", italic=True, space_after=4)
    for step in [
        "Ingestion — BASIL files on disk, or live coverage assembled from GDELT",
        "Preprocessing — sentence segmentation (spaCy)",
        "NER + cross-source canonicalization — align entity name variants into one identity",
        "Sentiment — per-entity tone, per source (RoBERTa)",
        "Framing — embedding-based sentence alignment and divergence scoring (MiniLM)",
        "Omission — entity-importance gaps and KeyBERT topic-overlap gaps",
        "Narrative clustering — KMeans over sentence embeddings",
        "Storage + API — PostgreSQL behind FastAPI GET endpoints",
        "Frontend — React panels rendering each signal",
    ]:
        bullet(doc, step)

    add_heading(doc, "2.1 Models", 2)
    body(doc, "Four of the five models are pretrained and run on every event with no training "
              "step; the fifth (a sentence-level bias classifier) is optional and fine-tuned "
              "separately on Kaggle.")
    styled_table(doc,
        ["Model", "Role", "Runs by default", "Trained by us"],
        [
            ["spaCy en_core_web_sm", "Named-entity recognition", "Yes", "No — pretrained"],
            ["cardiffnlp RoBERTa", "Entity sentiment", "Yes", "No — pretrained"],
            ["all-MiniLM-L6-v2", "Framing alignment", "Yes", "No — pretrained"],
            ["KeyBERT", "Topic-overlap omission", "Yes", "No — pretrained"],
            ["BERT / RoBERTa / DistilBERT", "Sentence bias classifier", "Optional", "Yes — Kaggle fine-tune"],
        ])

    # ---------- 3. Analysis methods ----------
    add_heading(doc, "3. Analysis Methods", 1)

    add_heading(doc, "3.1 Cross-source entity canonicalization", 2)
    body(doc,
         "Sources spell the same real-world entity differently — 'Trump', 'Donald Trump', "
         "'President Trump'. Left unhandled, these split one entity into several rows and corrupt "
         "every downstream per-source comparison. NewsLens pools entities from all sources, "
         "canonicalizes them together using fuzzy string matching with a subset-matching fallback "
         "for people's names, then maps the resolved names back to each source. This yields "
         "accurate per-source frequencies and omission percentages.")

    add_heading(doc, "3.2 Framing score", 2)
    body(doc,
         "Naively pairing sentences by shared keywords measures topic overlap, not framing. "
         "Instead, NewsLens embeds every sentence, aligns each to its most similar counterpart in "
         "the other source, keeps only alignments above a similarity floor, and scores divergence "
         "as the mean of (1 − similarity) over those kept pairs. Because each sentence is matched "
         "to its single most-similar counterpart, the score is a lower bound on framing "
         "divergence — 'how different are these sources even in their most similar description', "
         "not a complete framing measure. This caveat is stated in the code, the report, and the "
         "user interface.")

    add_heading(doc, "3.3 Omission (two signals)", 2)
    body(doc,
         "Entity-based omission weights each entity by the fraction of sources that mention it, "
         "so failing to mention something every other source treats as central counts more than "
         "skipping a passing detail. A source's omission score is the weighted fraction of that "
         "reference importance it fails to cover. A second, KeyBERT-based topic-overlap signal "
         "catches omitted themes that named-entity omission misses — a source that never engages "
         "a topic even though the same entities appear.")

    # ---------- 4. Live ingestion ----------
    add_heading(doc, "4. Extension 1 — Live News Ingestion", 1)
    body(doc,
         "The original system worked only on the ~300 pre-grouped articles in the BASIL dataset. "
         "The live-ingestion extension lets a user analyze any current story. The downstream "
         "analysis is unchanged — the only new problem is assembling a comparable, multi-source "
         "event on demand.")
    for step in [
        "Query — the user's topic is sent to the open GDELT Doc 2.0 API (no API key needed), "
        "as a stopword-filtered keyword-AND search so results stay focused on one story.",
        "One article per outlet — because the goal is to compare outlets, not individual stories, "
        "and the pipeline keys off a unique source label per article.",
        "Body extraction — each article's HTML is reduced to clean body text with readability-lxml.",
        "Wire-copy dedup — near-identical syndicated copy (AP/Reuters) is detected via MiniLM "
        "similarity and dropped, so framing scores reflect real editorial difference.",
        "Lean tagging — each source is mapped to a coarse political-lean bucket (left / center / "
        "right) from a curated domain registry, shown as a badge in the interface.",
        "Caching — results are cached per query for six hours so repeat searches are instant.",
    ]:
        bullet(doc, step)
    body(doc,
         "Known constraint: GDELT's free tier rate-limits to roughly one request every five "
         "seconds and can be slow to respond. The loader retries with backoff and reports these "
         "conditions clearly, but a live demo should not depend on the external API responding "
         "instantly. Lean labels are an orientation aid, not a ground-truth rating of any outlet.",
         italic=True)

    # ---------- 5. Trust & validation ----------
    add_heading(doc, "5. Extension 2 — Trust and Validation", 1)
    body(doc,
         "A public bias tool that presents unvalidated scores as fact can do real harm. This "
         "extension asks whether the framing metric actually tracks anything real, and reshapes "
         "the interface around the honest answer.")

    add_heading(doc, "5.1 Method", 2)
    body(doc,
         "The one ground truth available is BASIL's human bias-span annotations. For every event "
         "and every source pair, the real framing alignment is run; each aligned sentence pair is "
         "labeled 'human-biased' if either sentence contains an annotated bias span. The question "
         "is whether the pairs the metric flags as more divergent are the ones humans flagged.")

    add_heading(doc, "5.2 Results", 2)
    body(doc, "Evaluated across all 100 BASIL events (4,311 aligned sentence pairs, 1,498 bias spans):")
    styled_table(doc,
        ["Measure", "Value", "Interpretation"],
        [
            ["ROC AUC (divergence → human-biased)", "0.558", "Weak — leans right, close to chance"],
            ["Point-biserial correlation", "0.098", "Small positive association"],
            ["Mean divergence, human-biased pairs", "0.375", "Higher than non-biased"],
            ["Mean divergence, other pairs", "0.351", "Baseline"],
        ])
    body(doc,
         "Reading: divergent phrasing does lean toward human-perceived bias, but only weakly. The "
         "honest conclusion is that the framing number should be read as 'these sources phrase "
         "things differently', never as 'this source is biased'. This result is reproducible via "
         "scripts/evaluate_framing.py and is stated directly in the application's interface.")

    add_heading(doc, "5.3 Interface changes", 2)
    bullet(doc, "A persistent 'How to read this' banner frames every panel as a signal of difference, not a verdict, and states the measured AUC.", "Trust banner: ")
    bullet(doc, "Panel descriptions were rewritten (sentiment is 'a model guess'; omission is 'relative to the sources shown, not suppression').", "Reframed copy: ")
    bullet(doc, "A source-count line and per-pair framing confidence (flagging scores backed by fewer than five aligned pairs) surface how trustworthy each number is.", "Confidence signals: ")

    # ---------- 6. Worked example ----------
    add_heading(doc, "6. Worked Example", 1)
    src = ", ".join(ex["sources"])
    bt = ex["by_type"]
    body(doc,
         f"A representative BASIL event compares {len(ex['sources'])} sources ({src}) covering the "
         f"2018 U.S. withdrawal of troops from Syria and the related resignation of envoy Brett "
         f"McGurk. The pipeline extracted {ex['n_entities']} canonicalized entity rows across the "
         f"sources ({bt.get('PERSON',0)} PERSON, {bt.get('ORG',0)} ORG, {bt.get('GPE',0)} GPE).")
    body(doc, "Framing divergence (lower bound) between each source pair:", space_after=4)
    styled_table(doc,
        ["Source pair", "Divergence", "Aligned pairs"],
        [[f"{a} vs {b}", f"{s}%", n] for (a, b, s, n) in ex["framing"]])
    body(doc, "Entity omission — share of event-wide entity importance each source does not mention:",
         space_after=4)
    styled_table(doc,
        ["Source", "Omission"],
        [[s, f"{v}%"] for s, v in ex["omission"].items()])
    body(doc,
         f"The most divergent aligned sentence pair in the {ex['framing'][0][0]} vs "
         f"{ex['framing'][0][1]} comparison (similarity {ex['pair']['sim']:.2f}) illustrates the "
         f"signal concretely:", space_after=4)
    q1 = doc.add_paragraph()
    q1.paragraph_format.left_indent = Inches(0.4)
    r = q1.add_run(f"[{ex['framing'][0][0]}]  {ex['pair']['a']}")
    r.italic = True
    r.font.color.rgb = BLUE
    q2 = doc.add_paragraph()
    q2.paragraph_format.left_indent = Inches(0.4)
    r = q2.add_run(f"[{ex['framing'][0][1]}]  {ex['pair']['b']}")
    r.italic = True
    r.font.color.rgb = RED
    body(doc,
         "Both sentences describe the same withdrawal, but one foregrounds an official's assessment "
         "of ISIS while the other foregrounds Republican concern over Russia and Iran — a visible "
         "difference in emphasis, which is exactly what the metric is designed to surface (and "
         "exactly the kind of difference it should not be over-read as a bias verdict).",
         space_after=8)
    if not ex["live"]:
        body(doc, "(Figures above are recorded values; the API was not running at generation time.)",
             italic=True, color=SOFT, size=9)

    # ---------- 7. Implementation ----------
    add_heading(doc, "7. Implementation", 1)
    styled_table(doc,
        ["Layer", "Technology"],
        [
            ["Backend API", "FastAPI + Uvicorn (Python 3.11)"],
            ["NLP / ML", "spaCy, Transformers, sentence-transformers, KeyBERT, scikit-learn, PyTorch"],
            ["Live ingestion", "GDELT Doc 2.0 API, requests, readability-lxml"],
            ["Storage", "PostgreSQL (psycopg 3)"],
            ["Frontend", "React + Vite"],
            ["Bias-classifier training", "Kaggle notebook (T4 GPU), BERT / RoBERTa / DistilBERT"],
        ])

    # ---------- 8. Limitations ----------
    add_heading(doc, "8. Limitations", 1)
    bullet(doc, "The framing metric is a lower bound and only weakly correlated with human bias labels (Section 5). It is descriptive, not a verdict.", "Framing: ")
    bullet(doc, "Omission is measured only relative to the specific sources shown; 'missing' does not imply deliberate suppression.", "Omission scope: ")
    bullet(doc, "Sentiment is an automated estimate and can misread quotes, sarcasm, and context.", "Sentiment: ")
    bullet(doc, "Live results depend on GDELT's free tier, which rate-limits and can return thin coverage for narrow or opinion-framed queries.", "Live coverage: ")
    bullet(doc, "Political-lean labels are a coarse orientation aid from a curated list, not a ground-truth rating.", "Lean labels: ")
    bullet(doc, "The optional bias classifier's training labels are approximate (substring matching rather than exact character offsets).", "Bias labels: ")

    # ---------- 9. Future work ----------
    add_heading(doc, "9. Future Work", 1)
    bullet(doc, "Deployment — containerize the backend, use managed PostgreSQL, host the frontend, and add rate-limiting so the tool is reachable at a public URL.", "Phase 3: ")
    bullet(doc, "Additional live sources — NewsAPI and Google News RSS to broaden coverage beyond GDELT.")
    bullet(doc, "Further validation — automated omission labeling and an adjusted Rand index for the narrative clustering.")
    bullet(doc, "Calibration — tune the framing similarity threshold against the validation set to improve the weak correlation.")

    # ---------- 10. How to run ----------
    add_heading(doc, "10. How to Run", 1)
    for step in [
        "Start PostgreSQL (Docker: docker compose up -d, or a local server on port 5432).",
        "Backend: activate the virtualenv and run  uvicorn api.main:app --reload --port 8000  "
        "(Swagger docs at /docs).",
        "Frontend: cd frontend && npm install && npm run dev  (opens at http://localhost:5173).",
        "In the app, open one of the pre-analyzed BASIL events, or type a live topic to analyze "
        "current coverage.",
        "To reproduce the framing validation:  python -m scripts.evaluate_framing",
    ]:
        bullet(doc, step)

    # ---------- 11. Conclusion ----------
    add_heading(doc, "11. Conclusion", 1)
    body(doc,
         "NewsLens delivers a complete, working NLP pipeline for comparative media analysis and "
         "extends it in the two directions that matter most for real-world usefulness: it can now "
         "analyze live news on any current topic, and it presents its results honestly, backed by "
         "an explicit validation of its central metric. The system's most important design "
         "commitment is restraint — it shows readers how sources differ and equips them to judge, "
         "rather than claiming to have judged for them. The measured, modest correlation between "
         "its framing signal and human bias labels is reported openly rather than hidden, which is "
         "precisely what a trustworthy analysis tool should do.")

    out = "NewsLens_Project_Report.docx"
    doc.save(out)
    print(f"Wrote {out}  (worked example from {'LIVE API' if ex['live'] else 'recorded fallback'})")


if __name__ == "__main__":
    build()
